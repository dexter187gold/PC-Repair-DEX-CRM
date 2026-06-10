import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import json
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches
import shutil

# Data storage
DATA_FILE = "dex_crm_data.json"
TEMPLATES = {
    "Small": "PC_Repair_DEX_Small_Business_Report_Template.docx",
    "Medium": "PC_Repair_DEX_Medium_Business_Report_Template.docx",
    "Enterprise": "PC_Repair_DEX_Enterprise_Report_Template.docx"
}

class DexCrmApp:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("PC REPAIR DEX CRM V1")
        self.root.geometry("1200x800")
        
        self.load_data()
        self.current_theme = "light"
        self.create_gui()
        
    def load_data(self):
        if os.path.exists(DATA_FILE):
            with open(DATA_FILE, 'r') as f:
                self.data = json.load(f)
        else:
            self.data = {
                "clients": [],
                "last_job_numbers": {"S": 0, "M": 0, "E": 0}
            }
    
    def save_data(self):
        with open(DATA_FILE, 'w') as f:
            json.dump(self.data, f, indent=4)
    
    def create_gui(self):
        # Menu for theme
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)
        theme_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="Theme", menu=theme_menu)
        theme_menu.add_command(label="Light Mode", command=self.light_mode)
        theme_menu.add_command(label="Dark Mode", command=self.dark_mode)
        
        # Main frame
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Job Type Selection
        tk.Label(main_frame, text="Job Tier:").pack(anchor="w")
        self.tier_var = tk.StringVar(value="Enterprise")
        ttk.Radiobutton(main_frame, text="Small", variable=self.tier_var, value="Small").pack(anchor="w")
        ttk.Radiobutton(main_frame, text="Medium", variable=self.tier_var, value="Medium").pack(anchor="w")
        ttk.Radiobutton(main_frame, text="Enterprise (Checkers)", variable=self.tier_var, value="Enterprise").pack(anchor="w")
        
        # Client Selection / Add
        tk.Label(main_frame, text="Client:").pack(anchor="w", pady=(10,0))
        self.client_frame = ttk.Frame(main_frame)
        self.client_frame.pack(fill="x")
        
        self.client_var = tk.StringVar()
        self.client_combo = ttk.Combobox(self.client_frame, textvariable=self.client_var)
        self.client_combo['values'] = [c['name'] for c in self.data['clients']]
        self.client_combo.pack(side="left", fill="x", expand=True)
        
        ttk.Button(self.client_frame, text="Add New Client", command=self.add_client).pack(side="right")
        
        # Job Details
        details_frame = ttk.LabelFrame(main_frame, text="Job Details", padding="10")
        details_frame.pack(fill="x", pady=10)
        
        # Technician
        tk.Label(details_frame, text="Technician Name:").grid(row=0, column=0, sticky="w")
        self.tech_var = tk.StringVar(value="Your Name")
        ttk.Entry(details_frame, textvariable=self.tech_var).grid(row=0, column=1, sticky="ew")
        
        # Date
        tk.Label(details_frame, text="Date:").grid(row=1, column=0, sticky="w")
        self.date_var = tk.StringVar(value=datetime.now().strftime("%d/%m/%Y"))
        ttk.Entry(details_frame, textvariable=self.date_var).grid(row=1, column=1, sticky="ew")
        
        # Start/End Time
        tk.Label(details_frame, text="Start Time:").grid(row=2, column=0, sticky="w")
        self.start_time = tk.StringVar()
        ttk.Entry(details_frame, textvariable=self.start_time).grid(row=2, column=1, sticky="ew")
        
        tk.Label(details_frame, text="End Time:").grid(row=3, column=0, sticky="w")
        self.end_time = tk.StringVar()
        ttk.Entry(details_frame, textvariable=self.end_time).grid(row=3, column=1, sticky="ew")
        
        # Issue
        tk.Label(details_frame, text="Issue Reported:").grid(row=4, column=0, sticky="w")
        self.issue_text = tk.Text(details_frame, height=3, width=50)
        self.issue_text.grid(row=4, column=1, sticky="ew")
        
        details_frame.columnconfigure(1, weight=1)
        
        # Generate Button
        ttk.Button(main_frame, text="Generate Report", command=self.generate_report).pack(pady=20)
        
        # Status
        self.status_label = tk.Label(main_frame, text="", fg="green")
        self.status_label.pack()
    
    def add_client(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("Add New Client")
        dialog.geometry("400x300")
        
        tk.Label(dialog, text="Business Name:").pack(pady=5)
        name_entry = ttk.Entry(dialog, width=40)
        name_entry.pack()
        
        tk.Label(dialog, text="Contact Person:").pack(pady=5)
        contact_entry = ttk.Entry(dialog, width=40)
        contact_entry.pack()
        
        tk.Label(dialog, text="Phone:").pack(pady=5)
        phone_entry = ttk.Entry(dialog, width=40)
        phone_entry.pack()
        
        tk.Label(dialog, text="Email:").pack(pady=5)
        email_entry = ttk.Entry(dialog, width=40)
        email_entry.pack()
        
        def save_client():
            if name_entry.get().strip():
                new_client = {
                    "name": name_entry.get().strip(),
                    "contact": contact_entry.get().strip(),
                    "phone": phone_entry.get().strip(),
                    "email": email_entry.get().strip()
                }
                self.data["clients"].append(new_client)
                self.save_data()
                self.client_combo['values'] = [c['name'] for c in self.data['clients']]
                self.client_var.set(new_client['name'])
                dialog.destroy()
                messagebox.showinfo("Success", "Client added!")
        
        ttk.Button(dialog, text="Save Client", command=save_client).pack(pady=10)
    
    def light_mode(self):
        self.current_theme = "light"
        self.root.configure(bg="white")
        # Simple theme switch - in full version use ttk themes
    
    def dark_mode(self):
        self.current_theme = "dark"
        self.root.configure(bg="#2e2e2e")
        # More styling can be added
    
    def generate_report(self):
        tier = self.tier_var.get()
        client_name = self.client_var.get()
        
        if not client_name:
            messagebox.showerror("Error", "Please select or add a client")
            return
        
        # Generate Job ID
        prefix = tier[0]
        year = datetime.now().year
        month = f"{datetime.now().month:02d}"
        seq = self.data["last_job_numbers"][prefix] + 1
        job_id = f"DEX-{prefix}-{year}-{month}-{seq:04d}"
        self.data["last_job_numbers"][prefix] = seq
        self.save_data()
        
        template_path = TEMPLATES[tier]
        if not os.path.exists(template_path):
            messagebox.showerror("Error", f"Template {template_path} not found!")
            return
        
        try:
            doc = Document(template_path)
            
            # Replace basic placeholders (expand as needed)
            for paragraph in doc.paragraphs:
                if "[Your Name]" in paragraph.text:
                    paragraph.text = paragraph.text.replace("[Your Name]", self.tech_var.get())
                if "[DD/MM/YYYY]" in paragraph.text:
                    paragraph.text = paragraph.text.replace("[DD/MM/YYYY]", self.date_var.get())
                if "DEX-" in paragraph.text or "Report ID" in paragraph.text:
                    paragraph.text = paragraph.text.replace("DEX-[Tier]", job_id)
            
            # More advanced replacements can be added based on full template content
            
            output_dir = "Generated_Reports"
            os.makedirs(output_dir, exist_ok=True)
            output_path = os.path.join(output_dir, f"{job_id}_{client_name.replace(' ', '_')}.docx")
            doc.save(output_path)
            
            self.status_label.config(text=f"Report generated: {output_path}")
            messagebox.showinfo("Success", f"Report saved as {output_path}\nJob ID: {job_id}")
            
            # Optional: open the file
            if os.name == 'nt':
                os.startfile(output_path)
                
        except Exception as e:
            messagebox.showerror("Error", f"Failed to generate report: {str(e)}")
    
    def run(self):
        self.root.mainloop()

if __name__ == "__main__":
    app = DexCrmApp()
    app.run()
